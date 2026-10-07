"""OpenAI를 호출하지 않고 CSV, 비율, 한글 그림, 워드 파일만 확인한다."""

from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import app  # noqa: E402


def sample_frame():
    import pandas as pd

    return pd.DataFrame(
        [
            ["1", "2026-01-02", "김민수", "반도체", "경기", "세정기", "품질", "배송 지연과 포장 파손이 반복됩니다"],
            ["2", "2026-01-03", "이서연", "반도체", "서울", "세정기", "납기", "배송 지연이 가장 큽니다"],
            ["3", "2026-01-04", "박지훈", "자동차", "부산", "도장기", "품질", "도장 불량과 색상 차이가 있습니다"],
            ["4", "2026-01-05", "최하늘", "자동차", "대구", "도장기", "서비스", "응대 지연과 설명 부족이 반복됩니다"],
        ],
        columns=list(app.REQUIRED_COLUMNS),
    )


def write_csv(frame, encoding: str) -> Path:
    folder = tempfile.TemporaryDirectory()
    path = Path(folder.name) / "voc.csv"
    frame.to_csv(path, index=False, encoding=encoding)
    write_csv.folders.append(folder)
    return path


write_csv.folders = []


class CsvTests(unittest.TestCase):
    def test_reads_utf8_sig_and_cp949(self):
        frame = sample_frame()
        for encoding in ("utf-8-sig", "cp949"):
            loaded, used = app.read_voc_csv(write_csv(frame, encoding))
            self.assertEqual(used, encoding)
            self.assertEqual(loaded.iloc[0]["산업군"], "반도체")
            self.assertEqual(list(loaded.columns[:8]), list(app.REQUIRED_COLUMNS))

    def test_missing_column_is_rejected(self):
        frame = sample_frame().drop(columns=["불만"])
        with self.assertRaises(app.VocError):
            app.read_voc_csv(write_csv(frame, "utf-8-sig"))

    def test_empty_file_is_rejected(self):
        import pandas as pd

        frame = pd.DataFrame(columns=list(app.REQUIRED_COLUMNS))
        with self.assertRaises(app.VocError):
            app.read_voc_csv(write_csv(frame, "utf-8-sig"))


class StatsTests(unittest.TestCase):
    def test_ratios_keep_two_decimals(self):
        table = app.ratio_table(sample_frame(), "분야")
        ratios = {row["분야"]: row["비율(%)"] for row in table.to_dict("records")}
        self.assertEqual(ratios["품질"], 50.00)
        self.assertEqual(ratios["납기"], 25.00)
        self.assertEqual(ratios["서비스"], 25.00)
        self.assertAlmostEqual(sum(ratios.values()), 100.0, places=2)

    def test_chart_uses_korean_font_and_percent_text(self):
        table = app.ratio_table(sample_frame(), "분야")
        figure = app.build_ratio_chart(table, "분야")
        self.assertIn("Malgun Gothic", figure.layout.font.family)
        self.assertEqual(figure.layout.yaxis.tickformat, ".2f")
        self.assertCountEqual(list(figure.data[0].text), ["50.00%", "25.00%", "25.00%"])

    def test_chart_png_is_written(self):
        table = app.ratio_table(sample_frame(), "산업군")
        figure = app.build_ratio_chart(table, "산업군")
        with tempfile.TemporaryDirectory() as folder:
            dest = Path(folder) / "chart.png"
            app.save_chart_image(figure, dest)
            self.assertGreater(dest.stat().st_size, 5_000)

    def test_stats_brief_keeps_two_decimals_and_drops_braces(self):
        frame = sample_frame()
        frame.loc[0, "불만"] = "밸브{압력} 이상"
        text = app.complaint_brief(frame)
        self.assertNotIn("{", text)
        self.assertIn("밸브(압력)", text)
        brief = app.stats_brief(sample_frame())
        self.assertIn("50.00%", brief)
        self.assertIn("25.00%", brief)


class KeywordTests(unittest.TestCase):
    def test_korean_keywords_drop_particles_and_stopwords(self):
        counts = app.extract_keywords(app.complaint_texts(sample_frame()))
        self.assertEqual(counts["배송"], 2)
        self.assertEqual(counts["지연"], 3)
        self.assertIn("포장", counts)
        self.assertIn("파손", counts)
        self.assertIn("도장", counts)
        self.assertIn("불량", counts)
        self.assertNotIn("있습니다", counts)
        self.assertNotIn("큽니다", counts)

    def test_wordcloud_is_not_blank(self):
        image = app.build_wordcloud(sample_frame())
        self.assertEqual(image.size, (960, 540))
        low, high = image.convert("L").getextrema()
        self.assertNotEqual(low, high)


class ReportFileTests(unittest.TestCase):
    def test_docx_keeps_korean_text_and_font(self):
        content = app.ReportContent(
            overview="반도체와 자동차 비중이 같습니다.",
            issues="배송 지연이 반도체 쪽에서 반복됩니다.",
            actions="출고 점검 주기를 줄입니다.",
            issue_rows=[("반도체", "배송 지연", "출고 점검"), ("자동차", "도장 불량", "색상 기준 정비")],
        )
        with tempfile.TemporaryDirectory() as folder:
            dest = Path(folder) / "report.docx"
            app.build_report_docx(sample_frame(), content, dest)
            document = Document(dest)
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            self.assertIn("고객사 VOC 분석 보고서", text)
            self.assertIn("반도체와 자동차 비중이 같습니다.", text)
            self.assertIn("주요 이슈", text)
            self.assertIn("대응 방안", text)
            self.assertGreaterEqual(len(document.inline_shapes), 3)

            style_fonts = document.styles["Normal"].element.find(qn("w:rPr")).find(qn("w:rFonts"))
            self.assertEqual(style_fonts.get(qn("w:eastAsia")), "맑은 고딕")
            found_run = False
            for paragraph in document.paragraphs:
                for run in paragraph.runs:
                    if "반도체" not in run.text:
                        continue
                    fonts = run._element.find(qn("w:rPr")).find(qn("w:rFonts"))
                    self.assertEqual(fonts.get(qn("w:eastAsia")), "맑은 고딕")
                    found_run = True
            self.assertTrue(found_run)


class HandlerTests(unittest.TestCase):
    def test_upload_and_stats_handlers(self):
        path = write_csv(sample_frame(), "utf-8-sig")
        uploaded = list(app.on_upload(str(path), None))
        frame, _preview, upload_log = uploaded[-1]
        self.assertEqual(len(frame), 4)
        self.assertIn("[파일 업로드]", upload_log)

        analyzed = list(app.on_stats("제품명", frame))
        stats_log, figure, table = analyzed[-1]
        self.assertIn("[VOC Agent]", stats_log)
        self.assertIn("50.00%", list(figure.data[0].text))
        self.assertIn("50.00", table["비율(%)"].tolist())

    def test_wordcloud_handler_returns_image(self):
        result = list(app.on_wordcloud(sample_frame()))
        log, image, table = result[-1]
        self.assertIn("워드클라우드", log)
        self.assertEqual(image.size, (960, 540))
        self.assertIn("배송", table["키워드"].tolist())

    def test_report_without_file_does_not_call_model(self):
        result = list(app.on_report(None))
        log, report_path = result[-1]
        self.assertIn("업로드", log)
        self.assertIsNone(report_path)

    def test_ui_builds(self):
        self.assertIsNotNone(app.build_ui())

    def test_agent_order(self):
        self.assertEqual([item["key"] for item in app.AGENT_SPECS], ["voc", "issue", "report"])


if __name__ == "__main__":
    unittest.main()
