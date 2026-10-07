"""뉴스 수집·분석·보고서. LangGraph 노드 news, analyze, report."""

from __future__ import annotations

import html
import os
import re
import shutil
from datetime import datetime
from functools import lru_cache
from pathlib import Path
from typing import Callable, TypedDict

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt
from dotenv import load_dotenv
from langchain_core.messages import HumanMessage, SystemMessage
from langchain_openai import ChatOpenAI
from langgraph.graph import END, START, StateGraph
from PIL import Image
from plotly import graph_objects as go
from pydantic import BaseModel, Field
from wordcloud import WordCloud


BASE_DIR = Path(__file__).resolve().parent
load_dotenv(BASE_DIR / ".env")

OUTPUT_DIR = BASE_DIR / "outputs"
MODEL_NAME = os.environ.get("OPENAI_MODEL", "gpt-4o-mini")
FONT_NAME = "맑은 고딕"
FONT_FAMILY = "Malgun Gothic"
MAX_COUNT = 30
MAX_PAGES = 4
SUMMARY_LIMIT = 1000
CHART_LIMIT = 15
CLOUD_LIMIT = 80

SOURCES = {"구글 뉴스": "google", "네이버 뉴스": "naver"}
SOURCE_LABELS = {value: key for key, value in SOURCES.items()}

SearchFn = Callable[[dict], dict]

_PARTICLES = tuple(
    sorted(
        {
            "으로서",
            "로서",
            "에서는",
            "으로는",
            "에게는",
            "한테는",
            "에서",
            "으로",
            "에게",
            "한테",
            "까지",
            "부터",
            "처럼",
            "보다",
            "은",
            "는",
            "이",
            "가",
            "을",
            "를",
            "의",
            "에",
            "와",
            "과",
            "도",
            "만",
            "로",
        },
        key=len,
        reverse=True,
    )
)
_ENDINGS = tuple(
    sorted(
        {
            "되었습니다",
            "했습니다",
            "드립니다",
            "바랍니다",
            "습니다",
            "입니다",
            "합니다",
            "됩니다",
            "했어요",
            "네요",
            "는데",
            "어서",
            "아서",
            "해서",
            "하고",
            "이며",
            "으며",
            "거나",
            "지만",
            "니다",
        },
        key=len,
        reverse=True,
    )
)
_STOPWORDS = {
    "그리고",
    "그러나",
    "하지만",
    "매우",
    "너무",
    "정말",
    "가장",
    "관련",
    "대한",
    "대해",
    "위해",
    "위한",
    "통해",
    "있는",
    "없는",
    "하는",
    "되는",
    "및",
    "또는",
    "또한",
    "이번",
    "오늘",
    "어제",
    "지난",
    "저희",
    "우리",
    "때문",
    "때문에",
    "이후",
    "이전",
    "현재",
    "그냥",
    "조금",
    "약간",
    "계속",
    "다시",
    "아직",
    "이미",
    "더욱",
    "있습니다",
    "없습니다",
    "합니다",
    "됩니다",
    "했습니다",
    "입니다",
    "기자",
    "뉴스",
    "단독",
    "속보",
    "종합",
    "사진",
    "영상",
    "출처",
    "기사",
    "관련기사",
    "the",
    "and",
    "for",
    "with",
    "from",
    "that",
    "this",
    "news",
}

SYSTEM_PROMPT = (
    "너는 한국어 뉴스 분석가다. "
    "입력된 기사에 적힌 내용만 사용하고, 없는 수치·인용·사건을 만들지 않는다. "
    "기사를 주제별로 묶고, 각 묶음에 제목, 1000자 이내 요약, 시사점을 작성한다. "
    "요약과 시사점은 한국어 존댓말이다. "
    "기사 순번은 입력에 있는 번호만 쓰고, 한 기사는 한 그룹에만 넣는다. "
    "그룹은 기사 수에 맞게 1개에서 4개로 나눈다. "
    "빈도 키워드는 참고만 하고, 그룹 제목은 기사 내용으로 정한다."
)


class NewsError(ValueError):
    """사용자가 고칠 수 있는 입력·수집·분석 오류."""


class GroupDraft(BaseModel):
    title: str = Field(description="그룹 제목. 한국어")
    summary: str = Field(description="그룹 요약. 한국어 1000자 이내")
    insight: str = Field(description="주요 시사점. 한국어")
    article_numbers: list[int] = Field(description="이 그룹에 넣을 기사 순번")


class GroupDrafts(BaseModel):
    groups: list[GroupDraft] = Field(description="주제별 뉴스 그룹")


class NewsState(TypedDict, total=False):
    keyword: str
    count: int
    source: str
    work_dir: str
    articles: list[dict]
    notice: str
    groups: list[dict]
    keywords: list[dict]
    wordcloud_path: str
    chart_path: str
    report_path: str
    used_fallback: bool


def scrub_secrets(text: str) -> str:
    cleaned = re.sub(r"(api_key=)[^&\s]+", r"\1***", text or "", flags=re.I)
    return cleaned[:500]


def require_runtime_keys() -> None:
    if not (os.environ.get("SERPAPI_API_KEY") or os.environ.get("SERPAPI_KEY")):
        raise NewsError("SERPAPI_API_KEY가 없습니다. news_analysis/.env 파일에 키를 넣으세요.")
    if not os.environ.get("OPENAI_API_KEY"):
        raise NewsError("OPENAI_API_KEY 환경 변수가 없습니다.")


def validate_request(keyword: object, count: object, source: object) -> tuple[str, int, str]:
    text = "" if keyword is None else str(keyword).strip()
    if not text:
        raise NewsError("뉴스 키워드를 입력하세요.")
    if len(text) > 80:
        raise NewsError("키워드는 80자 이내로 입력하세요.")
    if isinstance(count, bool):
        raise NewsError("수집 건수는 숫자로 입력하세요.")
    try:
        number = float(count)
    except (TypeError, ValueError):
        raise NewsError("수집 건수는 숫자로 입력하세요.") from None
    if not number.is_integer():
        raise NewsError("수집 건수는 정수로 입력하세요.")
    size = int(number)
    if size < 1 or size > MAX_COUNT:
        raise NewsError(f"수집 건수는 1부터 {MAX_COUNT}까지 입력하세요.")
    source_text = "" if source is None else str(source).strip()
    if source_text in SOURCES:
        source_key = SOURCES[source_text]
    elif source_text in SOURCE_LABELS:
        source_key = source_text
    else:
        raise NewsError("뉴스는 구글 뉴스 또는 네이버 뉴스 중에서 선택하세요.")
    return text, size, source_key


def clean_text(value: object) -> str:
    if value is None:
        return ""
    text = html.unescape(str(value))
    text = re.sub(r"<[^>]+>", "", text)
    text = text.replace("\xa0", " ").replace("\u200b", "")
    text = re.sub(r"\s+", " ", text).strip()
    return text


def clamp_text(value: object, limit: int = SUMMARY_LIMIT) -> str:
    text = "" if value is None else str(value).replace("\r\n", "\n").strip()
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    if len(text) <= limit:
        return text
    return text[:limit].rstrip()


def format_date(value: object) -> str:
    text = clean_text(value)
    if not text:
        return "날짜 없음"
    iso_match = re.search(r"(\d{4})-(\d{2})-(\d{2})", text)
    if iso_match:
        return "-".join(iso_match.groups())
    head = text[:10]
    for pattern in ("%m/%d/%Y", "%Y.%m.%d", "%Y/%m/%d"):
        try:
            return datetime.strptime(head, pattern).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return text


def canonical_url(value: object) -> str:
    return clean_text(value).split("#", 1)[0].strip()


def _record(title: object, url: object, summary: object, date: object) -> dict | None:
    link = canonical_url(url)
    headline = clean_text(title)
    if not headline or not link.startswith(("http://", "https://")):
        return None
    body = clean_text(summary)
    return {
        "날짜": format_date(date),
        "제목": headline[:300],
        "주요내용": (body or "요약 없음")[:SUMMARY_LIMIT],
        "url": link,
    }


def _google_source_name(source: object) -> str:
    if isinstance(source, dict):
        return clean_text(source.get("name") or source.get("title"))
    return clean_text(source)


def parse_google_entry(item: dict) -> dict | None:
    if not isinstance(item, dict):
        return None
    date = item.get("iso_date") or item.get("published_at") or item.get("date")
    summary = item.get("snippet") or _google_source_name(item.get("source"))
    return _record(item.get("title"), item.get("link"), summary, date)


def iter_google_items(payload: dict) -> list[dict]:
    rows: list[dict] = []
    for item in payload.get("news_results") or []:
        if not isinstance(item, dict):
            continue
        if "highlight" in item or "stories" in item:
            highlight = item.get("highlight")
            if isinstance(highlight, dict):
                parsed = parse_google_entry(highlight)
                if parsed:
                    rows.append(parsed)
            for story in item.get("stories") or []:
                parsed = parse_google_entry(story)
                if parsed:
                    rows.append(parsed)
            continue
        parsed = parse_google_entry(item)
        if parsed:
            rows.append(parsed)
    return rows


def parse_naver_entry(item: dict) -> dict | None:
    if not isinstance(item, dict):
        return None
    info = item.get("news_info") if isinstance(item.get("news_info"), dict) else {}
    date = info.get("news_date") or item.get("date")
    return _record(item.get("title"), item.get("link"), item.get("snippet"), date)


def iter_naver_items(payload: dict) -> list[dict]:
    rows = []
    for item in payload.get("news_results") or []:
        parsed = parse_naver_entry(item)
        if parsed:
            rows.append(parsed)
    return rows


def number_articles(items: list[dict], count: int) -> tuple[list[dict], str]:
    seen: set[str] = set()
    selected: list[dict] = []
    for item in items:
        link = item["url"]
        if link in seen:
            continue
        seen.add(link)
        selected.append(item)
        if len(selected) >= count:
            break
    articles = [
        {
            "순번": index,
            "날짜": item["날짜"],
            "제목": item["제목"],
            "주요내용": item["주요내용"],
            "url": item["url"],
        }
        for index, item in enumerate(selected, 1)
    ]
    notice = ""
    if articles and len(articles) < count:
        notice = f"요청 {count}건 중 {len(articles)}건만 수집되었습니다."
    return articles, notice


def _payload_or_error(payload: dict) -> dict:
    if not isinstance(payload, dict):
        raise NewsError("SerpAPI 응답 형식이 올바르지 않습니다.")
    error = payload.get("error")
    if error:
        raise NewsError(f"SerpAPI 오류: {scrub_secrets(str(error))}")
    return payload


def default_search(params: dict) -> dict:
    api_key = os.environ.get("SERPAPI_API_KEY") or os.environ.get("SERPAPI_KEY")
    if not api_key:
        raise NewsError("SERPAPI_API_KEY가 없습니다. news_analysis/.env 파일에 키를 넣으세요.")
    try:
        import serpapi
    except ImportError as exc:
        raise NewsError("serpapi 패키지가 설치되어 있지 않습니다.") from exc
    try:
        result = serpapi.Client(api_key=api_key).search(params)
    except Exception as exc:
        detail = getattr(exc, "error", None) or str(exc)
        raise NewsError(f"SerpAPI 요청에 실패했습니다. {scrub_secrets(str(detail))}") from exc
    data = result.as_dict() if hasattr(result, "as_dict") else dict(result)
    return _payload_or_error(data)


def collect_news(keyword: str, count: int, source: str, search_fn: SearchFn | None = None) -> tuple[list[dict], str]:
    search = search_fn or default_search
    if source == "google":
        payload = _payload_or_error(
            search(
                {
                    "engine": "google_news",
                    "q": keyword,
                    "gl": "kr",
                    "hl": "ko",
                    "so": 1,
                }
            )
        )
        items = iter_google_items(payload)
    elif source == "naver":
        items = []
        seen: set[str] = set()
        for page in range(1, MAX_PAGES + 1):
            if len(items) >= count:
                break
            payload = _payload_or_error(
                search(
                    {
                        "engine": "naver",
                        "query": keyword,
                        "where": "news",
                        "sort_by": 1,
                        "page": page,
                    }
                )
            )
            added = 0
            for item in iter_naver_items(payload):
                if item["url"] in seen:
                    continue
                seen.add(item["url"])
                items.append(item)
                added += 1
                if len(items) >= count:
                    break
            if added == 0:
                break
    else:
        raise NewsError("뉴스는 구글 뉴스 또는 네이버 뉴스 중에서 선택하세요.")

    articles, notice = number_articles(items, count)
    if not articles:
        raise NewsError("수집된 뉴스가 없습니다. 키워드를 바꿔 다시 시도하세요.")
    return articles, notice


@lru_cache(maxsize=1)
def korean_font_path() -> str:
    candidates = (
        r"C:\Windows\Fonts\malgun.ttf",
        r"C:\Windows\Fonts\malgunsl.ttf",
        "/usr/share/fonts/truetype/nanum/NanumGothic.ttf",
        "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
    )
    for candidate in candidates:
        if Path(candidate).is_file():
            return candidate
    raise NewsError("한글 폰트를 찾지 못했습니다. Windows 맑은 고딕(malgun.ttf)이 필요합니다.")


def _stem(token: str) -> str:
    for ending in _ENDINGS:
        if token.endswith(ending) and len(token) - len(ending) >= 2:
            token = token[: -len(ending)]
            break
    for particle in _PARTICLES:
        if token.endswith(particle) and len(token) - len(particle) >= 2:
            token = token[: -len(particle)]
            break
    return token


def extract_keywords(texts: list[str]) -> dict[str, int]:
    counts: dict[str, int] = {}
    for text in texts:
        for token in re.findall(r"[A-Za-z]{2,}", text):
            word = token.lower()
            if word in _STOPWORDS:
                continue
            counts[word] = counts.get(word, 0) + 1
        for token in re.findall(r"[가-힣]+", text):
            if token in _STOPWORDS:
                continue
            word = _stem(token)
            if len(word) < 2 or word in _STOPWORDS:
                continue
            counts[word] = counts.get(word, 0) + 1
    return counts


def rank_keywords(articles: list[dict], limit: int = CLOUD_LIMIT) -> list[tuple[str, int]]:
    texts = [f"{article['제목']} {article['주요내용']}" for article in articles]
    ranked = sorted(extract_keywords(texts).items(), key=lambda item: (-item[1], item[0]))
    return ranked[:limit]


def keyword_rows(ranked: list[tuple[str, int]], limit: int = CHART_LIMIT) -> list[dict]:
    return [{"단어": word, "빈도": count} for word, count in ranked[:limit]]


def build_frequency_chart(rows: list[dict]) -> go.Figure:
    if not rows:
        raise NewsError("뉴스에서 키워드를 찾지 못했습니다.")
    ordered = list(reversed(rows[:CHART_LIMIT]))
    words = [row["단어"] for row in ordered]
    counts = [int(row["빈도"]) for row in ordered]
    figure = go.Figure(
        go.Bar(
            x=counts,
            y=words,
            orientation="h",
            text=[str(count) for count in counts],
            textposition="outside",
            textfont={"family": FONT_FAMILY, "size": 12},
            marker_color="#1B4F72",
            cliponaxis=False,
            hovertemplate="%{y}<br>빈도: %{x}회<extra></extra>",
        )
    )
    highest = max(counts) if counts else 1
    figure.update_layout(
        template="plotly_white",
        title="주요 단어 발생 빈도",
        font={"family": FONT_FAMILY, "size": 13},
        xaxis={"title": "발생 빈도", "range": [0, max(highest * 1.28, 1)]},
        yaxis={"title": "단어", "automargin": True},
        margin={"t": 70, "b": 50, "l": 20, "r": 40},
        showlegend=False,
        height=max(360, 42 * len(words) + 120),
    )
    return figure


def save_chart_image(figure: go.Figure, dest: Path) -> None:
    dest.parent.mkdir(parents=True, exist_ok=True)
    figure.write_image(str(dest), format="png", width=960, height=540, scale=1)


def build_wordcloud(ranked: list[tuple[str, int]]) -> Image.Image:
    if not ranked:
        raise NewsError("뉴스에서 키워드를 찾지 못했습니다.")
    cloud = WordCloud(
        font_path=korean_font_path(),
        width=960,
        height=540,
        background_color="white",
        colormap="Blues",
        prefer_horizontal=0.9,
        random_state=7,
    )
    return cloud.generate_from_frequencies(dict(ranked)).to_image()


def render_keyword_images(ranked: list[tuple[str, int]], work_dir: Path) -> tuple[str, str]:
    work_dir.mkdir(parents=True, exist_ok=True)
    chart_path = work_dir / "keyword_chart.png"
    cloud_path = work_dir / "wordcloud.png"
    save_chart_image(build_frequency_chart(keyword_rows(ranked)), chart_path)
    build_wordcloud(ranked).save(cloud_path, format="PNG")
    return str(chart_path), str(cloud_path)


def build_llm() -> ChatOpenAI:
    if not os.environ.get("OPENAI_API_KEY"):
        raise NewsError("OPENAI_API_KEY 환경 변수가 없습니다.")
    return ChatOpenAI(model=MODEL_NAME, temperature=0.2, timeout=120, max_retries=1)


def _drafts_from_result(result: object) -> list[GroupDraft]:
    if isinstance(result, GroupDrafts):
        return list(result.groups)
    if isinstance(result, dict):
        return [GroupDraft.model_validate(item) for item in result.get("groups") or []]
    groups = getattr(result, "groups", None)
    if groups is None:
        raise NewsError("뉴스 그룹 결과를 해석할 수 없습니다.")
    drafts = []
    for item in groups:
        drafts.append(item if isinstance(item, GroupDraft) else GroupDraft.model_validate(item))
    return drafts


def _article_refs(numbers: list[int], by_number: dict[int, dict]) -> list[dict]:
    refs = []
    for number in numbers:
        article = by_number[number]
        refs.append(
            {
                "순번": number,
                "날짜": article["날짜"],
                "제목": article["제목"],
                "url": article["url"],
            }
        )
    return refs


def fallback_group(articles: list[dict], title: str = "전체 뉴스") -> dict:
    parts = [f"{article['순번']}. {article['제목']} {article['주요내용']}" for article in articles]
    numbers = [int(article["순번"]) for article in articles]
    by_number = {int(article["순번"]): article for article in articles}
    return {
        "제목": title,
        "요약": clamp_text(" ".join(parts)) or "요약 없음",
        "시사점": "모델 응답에서 그룹을 만들지 못해, 수집 기사를 한 묶음으로 요약했습니다.",
        "기사번호": numbers,
        "기사": _article_refs(numbers, by_number),
    }


def normalize_groups(result: object, articles: list[dict]) -> tuple[list[dict], bool]:
    by_number = {int(article["순번"]): article for article in articles}
    used: set[int] = set()
    groups: list[dict] = []
    for draft in _drafts_from_result(result):
        numbers = []
        for number in draft.article_numbers:
            if number in by_number and number not in used:
                numbers.append(int(number))
                used.add(int(number))
        if not numbers:
            continue
        groups.append(
            {
                "제목": clean_text(draft.title) or "제목 없음",
                "요약": clamp_text(draft.summary) or "요약 없음",
                "시사점": clamp_text(draft.insight) or "시사점 없음",
                "기사번호": numbers,
                "기사": _article_refs(numbers, by_number),
            }
        )

    leftover = [article for article in articles if int(article["순번"]) not in used]
    used_fallback = False
    if not groups:
        groups = [fallback_group(articles)]
        used_fallback = True
    elif leftover:
        groups.append(fallback_group(leftover, "기타 기사"))
        used_fallback = True
    return groups, used_fallback


def group_news(keyword: str, articles: list[dict], keywords: list[dict], llm=None) -> tuple[list[dict], bool]:
    model = llm or build_llm()
    structured = model.with_structured_output(GroupDrafts)
    keyword_text = ", ".join(f"{row['단어']} {row['빈도']}" for row in keywords[:20]) or "없음"
    article_lines = []
    for article in articles:
        article_lines.append(
            f"{article['순번']}. [{article['날짜']}] {article['제목']}\n{clamp_text(article['주요내용'], 400)}\n{article['url']}"
        )
    message = (
        f"검색어: {keyword}\n"
        f"참고 키워드: {keyword_text}\n\n"
        "기사:\n"
        + "\n\n".join(article_lines)
    )
    result = structured.invoke(
        [
            SystemMessage(content=SYSTEM_PROMPT),
            HumanMessage(content=message),
        ]
    )
    return normalize_groups(result, articles)


def _set_style_font(style, font_name: str) -> None:
    style.font.name = font_name
    style_element = style.element
    r_pr = style_element.find(qn("w:rPr"))
    if r_pr is None:
        r_pr = OxmlElement("w:rPr")
        style_element.append(r_pr)
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        r_fonts.set(qn(attr), font_name)


def _add_run(paragraph, text: str, *, size: int = 11, bold: bool = False):
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = FONT_NAME
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        r_fonts.set(qn(attr), FONT_NAME)
    return run


def _add_paragraph(document: Document, text: str, *, size: int = 11, bold: bool = False, center: bool = False):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    if center:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_run(paragraph, text, size=size, bold=bold)
    return paragraph


def _add_lines(document: Document, text: str) -> None:
    lines = [line.strip() for line in str(text).splitlines() if line.strip()]
    if not lines:
        _add_paragraph(document, "내용 없음")
        return
    for line in lines:
        _add_paragraph(document, line)


def _write_cell(cell, text: str, *, bold: bool = False) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.clear()
    _add_run(paragraph, str(text), size=10, bold=bold)


def _add_table(document: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        _write_cell(table.rows[0].cells[index], header, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            _write_cell(cells[index], value)
    document.add_paragraph()


def safe_keyword(keyword: str) -> str:
    cleaned = re.sub(r'[\\/:*?"<>|]+', " ", keyword).strip()
    cleaned = re.sub(r"\s+", "_", cleaned)
    return cleaned[:40] or "news"


def build_report_docx(state: NewsState, dest: Path) -> Path:
    groups = state.get("groups") or []
    keywords = state.get("keywords") or []
    chart_path = Path(state.get("chart_path") or "")
    cloud_path = Path(state.get("wordcloud_path") or "")
    if not chart_path.is_file() or not cloud_path.is_file():
        raise NewsError("키워드 그래프 파일이 없습니다.")
    dest = Path(dest)
    dest.parent.mkdir(parents=True, exist_ok=True)

    document = Document()
    document.core_properties.title = "뉴스 분석 보고서"
    document.core_properties.language = "ko-KR"
    for style_name in ("Normal", "Title", "Heading 1", "Heading 2"):
        try:
            _set_style_font(document.styles[style_name], FONT_NAME)
        except KeyError:
            continue
    for section in document.sections:
        section.top_margin = Cm(1.8)
        section.bottom_margin = Cm(1.8)
        section.left_margin = Cm(1.8)
        section.right_margin = Cm(1.8)

    keyword = state.get("keyword") or ""
    source_label = SOURCE_LABELS.get(state.get("source", ""), state.get("source", ""))
    written_at = datetime.now().strftime("%Y-%m-%d %H:%M")
    _add_paragraph(document, "뉴스 분석 보고서", size=20, bold=True, center=True)
    _add_paragraph(
        document,
        f"검색어 {keyword} / 출처 {source_label} / 수집 {len(state.get('articles') or [])}건 / 작성 {written_at}",
    )

    for index, group in enumerate(groups, 1):
        _add_paragraph(document, f"{index}. {group['제목']}", size=16, bold=True)
        _add_paragraph(document, "요약", size=13, bold=True)
        _add_lines(document, group.get("요약") or "")
        _add_paragraph(document, "시사점", size=13, bold=True)
        _add_lines(document, group.get("시사점") or "")
        refs = group.get("기사") or []
        if refs:
            _add_table(
                document,
                ["순번", "날짜", "제목", "URL"],
                [[str(ref["순번"]), ref["날짜"], ref["제목"], ref["url"]] for ref in refs],
            )

    _add_paragraph(document, "키워드 그래프", size=16, bold=True)
    document.add_picture(str(chart_path), width=Inches(6.2))
    _add_paragraph(document, "워드클라우드", size=16, bold=True)
    document.add_picture(str(cloud_path), width=Inches(6.2))
    if keywords:
        _add_paragraph(document, "주요 단어 빈도", size=14, bold=True)
        _add_table(
            document,
            ["단어", "빈도"],
            [[row["단어"], str(row["빈도"])] for row in keywords],
        )

    document.save(dest)
    return dest


def copy_report(src: str | Path, folder: str) -> Path:
    source = Path(src)
    if not source.is_file():
        raise NewsError("저장할 보고서 파일이 없습니다. 분석을 먼저 실행하세요.")
    raw = str(folder or "").strip().strip('"')
    if not raw:
        raise NewsError("저장 폴더를 입력하세요.")
    dest_dir = Path(raw).expanduser()
    try:
        dest_dir.mkdir(parents=True, exist_ok=True)
    except OSError as exc:
        raise NewsError(f"저장 폴더를 만들 수 없습니다: {exc}") from exc
    if not dest_dir.is_dir():
        raise NewsError("저장 경로는 폴더여야 합니다.")
    dest = dest_dir / source.name
    try:
        shutil.copy2(source, dest)
    except OSError as exc:
        raise NewsError(f"보고서를 저장하지 못했습니다: {exc}") from exc
    return dest.resolve()


def build_graph(search_fn: SearchFn | None = None, llm=None):
    """news → analyze → report 순서로 뉴스를 모으고 보고서를 만든다."""

    def news_node(state: NewsState) -> dict:
        keyword, count, source = validate_request(state.get("keyword"), state.get("count"), state.get("source"))
        articles, notice = collect_news(keyword, count, source, search_fn=search_fn)
        return {"keyword": keyword, "count": count, "source": source, "articles": articles, "notice": notice}

    def analyze_node(state: NewsState) -> dict:
        articles = state.get("articles") or []
        if not articles:
            raise NewsError("분석할 뉴스가 없습니다.")
        ranked = rank_keywords(articles)
        if not ranked:
            raise NewsError("뉴스에서 키워드를 찾지 못했습니다.")
        work_dir = Path(state.get("work_dir") or OUTPUT_DIR)
        chart_path, cloud_path = render_keyword_images(ranked, work_dir)
        keywords = keyword_rows(ranked)
        groups, used_fallback = group_news(state.get("keyword") or "", articles, keywords, llm=llm)
        return {
            "groups": groups,
            "keywords": keywords,
            "chart_path": chart_path,
            "wordcloud_path": cloud_path,
            "used_fallback": used_fallback,
        }

    def report_node(state: NewsState) -> dict:
        work_dir = Path(state.get("work_dir") or OUTPUT_DIR)
        work_dir.mkdir(parents=True, exist_ok=True)
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        dest = work_dir / f"뉴스분석_{safe_keyword(state.get('keyword') or '')}_{stamp}.docx"
        build_report_docx(state, dest)
        return {"report_path": str(dest)}

    builder = StateGraph(NewsState)
    builder.add_node("news", news_node)
    builder.add_node("analyze", analyze_node)
    builder.add_node("report", report_node)
    builder.add_edge(START, "news")
    builder.add_edge("news", "analyze")
    builder.add_edge("analyze", "report")
    builder.add_edge("report", END)
    return builder.compile()
