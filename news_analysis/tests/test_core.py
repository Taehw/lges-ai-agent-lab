"""OpenAI와 SerpAPI를 호출하지 않고 수집 정리, 그래프, 워드 파일만 확인한다."""

from __future__ import annotations

import json
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import agent  # noqa: E402
import app  # noqa: E402


def google_payload() -> dict:
    return {
        "news_results": [
            {
                "highlight": {
                    "title": "반도체 수출 호조",
                    "link": "https://example.com/a",
                    "snippet": "반도체 수출과 메모리 수요가 늘었습니다",
                    "iso_date": "2026-03-01T01:00:00Z",
                },
                "stories": [
                    {
                        "title": "반도체 수출 호조 반복",
                        "link": "https://example.com/a#detail",
                        "snippet": "같은 기사",
                        "iso_date": "2026-03-01T02:00:00Z",
                    },
                    {
                        "title": "메모리 가격 상승",
                        "link": "https://example.com/b",
                        "snippet": "메모리 가격과 반도체 공급이 달라졌습니다",
                        "date": "03/02/2026, 09:00 AM",
                    },
                ],
            },
            {
                "title": "관세 정책 변화",
                "link": "https://example.com/c",
                "snippet": "미국 관세 정책이 반도체 공급망에 영향을 줍니다",
                "iso_date": "2026-03-03T00:00:00Z",
                "source": {"name": "예시뉴스"},
            },
        ]
    }


def naver_item(index: int, link: str | None = None) -> dict:
    return {
        "title": f"<b>반도체</b> 소식 {index}",
        "link": link or f"https://news.example/{index}",
        "snippet": "수출 호조 &amp; 메모리 수요",
        "news_info": {"news_date": "1시간 전", "press_name": "예시뉴스"},
    }


class RequestTests(unittest.TestCase):
    def test_accepts_source_label_and_count(self):
        keyword, count, source = agent.validate_request(" 반도체 ", 10, "구글 뉴스")
        self.assertEqual((keyword, count, source), ("반도체", 10, "google"))

    def test_rejects_blank_keyword_bad_count_and_source(self):
        with self.assertRaises(agent.NewsError):
            agent.validate_request("  ", 10, "구글 뉴스")
        with self.assertRaises(agent.NewsError):
            agent.validate_request("반도체", 0, "네이버 뉴스")
        with self.assertRaises(agent.NewsError):
            agent.validate_request("반도체", 10.5, "네이버 뉴스")
        with self.assertRaises(agent.NewsError):
            agent.validate_request("반도체", 31, "구글 뉴스")
        with self.assertRaises(agent.NewsError):
            agent.validate_request("반도체", 5, "다음 뉴스")

    def test_secret_is_removed_from_error_text(self):
        cleaned = agent.scrub_secrets("https://serpapi.com/search?api_key=secret123&q=1")
        self.assertNotIn("secret123", cleaned)
        self.assertIn("api_key=***", cleaned)


class CollectTests(unittest.TestCase):
    def test_google_flattens_stories_and_limits_count(self):
        captured = {}

        def search(params):
            captured.update(params)
            return google_payload()

        articles, notice = agent.collect_news("반도체", 2, "google", search_fn=search)
        self.assertEqual(captured["engine"], "google_news")
        self.assertEqual(captured["q"], "반도체")
        self.assertEqual(captured["gl"], "kr")
        self.assertEqual(captured["hl"], "ko")
        self.assertEqual(captured["so"], 1)
        self.assertEqual(len(articles), 2)
        self.assertEqual(articles[0]["날짜"], "2026-03-01")
        self.assertEqual(articles[1]["날짜"], "2026-03-02")
        self.assertEqual(list(articles[0].keys()), ["순번", "날짜", "제목", "주요내용", "url"])
        self.assertNotIn("secret", json.dumps(articles, ensure_ascii=False))
        self.assertIn("반도체", json.dumps(articles, ensure_ascii=False))
        self.assertEqual(notice, "")

    def test_naver_strips_html_and_stops_after_enough_pages(self):
        calls = []

        def search(params):
            calls.append(params["page"])
            if params["page"] == 1:
                return {"news_results": [naver_item(index) for index in range(1, 11)]}
            return {"news_results": [naver_item(index) for index in range(11, 13)]}

        articles, notice = agent.collect_news("반도체", 12, "naver", search_fn=search)
        self.assertEqual(calls, [1, 2])
        self.assertEqual(len(articles), 12)
        self.assertEqual(articles[0]["제목"], "반도체 소식 1")
        self.assertEqual(articles[0]["주요내용"], "수출 호조 & 메모리 수요")
        self.assertEqual(articles[0]["날짜"], "1시간 전")
        self.assertEqual(notice, "")

    def test_duplicate_page_stops_and_error_payload_is_rejected(self):
        calls = []

        def search(params):
            calls.append(params["page"])
            return {"news_results": [naver_item(1, "https://news.example/same")]}

        articles, notice = agent.collect_news("반도체", 5, "naver", search_fn=search)
        self.assertEqual(calls, [1, 2])
        self.assertEqual(len(articles), 1)
        self.assertIn("1건만", notice)

        def failed(_params):
            return {"error": "Invalid API key"}

        with self.assertRaises(agent.NewsError):
            agent.collect_news("반도체", 1, "google", search_fn=failed)


class KeywordTests(unittest.TestCase):
    def articles(self):
        return [
            {
                "순번": 1,
                "날짜": "2026-03-01",
                "제목": "반도체 수출 호조",
                "주요내용": "반도체 수출과 메모리 수요가 늘었습니다",
                "url": "https://example.com/a",
            },
            {
                "순번": 2,
                "날짜": "2026-03-02",
                "제목": "메모리 가격",
                "주요내용": "메모리 가격과 반도체 공급이 달라졌습니다",
                "url": "https://example.com/b",
            },
        ]

    def test_korean_keywords_drop_stopwords(self):
        counts = agent.extract_keywords(
            ["반도체 수출이 늘었습니다", "반도체 수출 호조와 메모리 수요가 있습니다"]
        )
        self.assertGreaterEqual(counts["반도체"], 2)
        self.assertGreaterEqual(counts["수출"], 2)
        self.assertIn("메모리", counts)
        self.assertNotIn("있습니다", counts)
        self.assertNotIn("뉴스", counts)

    def test_chart_uses_korean_font_and_frequency(self):
        rows = agent.keyword_rows(agent.rank_keywords(self.articles()))
        figure = agent.build_frequency_chart(rows)
        self.assertEqual(figure.layout.font.family, "Malgun Gothic")
        self.assertEqual(figure.layout.title.text, "주요 단어 발생 빈도")
        self.assertEqual(list(figure.data[0].y)[-1], rows[0]["단어"])
        self.assertEqual(list(figure.data[0].x)[-1], rows[0]["빈도"])

    def test_wordcloud_is_not_blank(self):
        image = agent.build_wordcloud(agent.rank_keywords(self.articles()))
        self.assertEqual(image.size, (960, 540))
        low, high = image.convert("L").getextrema()
        self.assertNotEqual(low, high)
        self.assertTrue(agent.korean_font_path().lower().endswith("malgun.ttf"))


class ReportTests(unittest.TestCase):
    def test_summary_is_clamped_and_docx_keeps_korean_font(self):
        articles = KeywordTests.articles(self)
        drafts = agent.GroupDrafts(
            groups=[
                agent.GroupDraft(
                    title="수출 호조",
                    summary="가" * 1200,
                    insight="메모리 수요가 함께 움직입니다.",
                    article_numbers=[1, 9],
                )
            ]
        )
        groups, used_fallback = agent.normalize_groups(drafts, articles)
        self.assertEqual(len(groups[0]["요약"]), 1000)
        self.assertEqual(groups[0]["기사번호"], [1])
        self.assertTrue(used_fallback)
        self.assertEqual(groups[1]["제목"], "기타 기사")

        ranked = agent.rank_keywords(articles)
        with tempfile.TemporaryDirectory() as folder:
            folder_path = Path(folder)
            chart_path, cloud_path = agent.render_keyword_images(ranked, folder_path)
            state = {
                "keyword": "반도체",
                "source": "google",
                "articles": articles,
                "groups": groups,
                "keywords": agent.keyword_rows(ranked),
                "chart_path": chart_path,
                "wordcloud_path": cloud_path,
            }
            dest = folder_path / "report.docx"
            agent.build_report_docx(state, dest)
            document = Document(dest)
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            self.assertIn("뉴스 분석 보고서", text)
            self.assertIn("수출 호조", text)
            self.assertIn("메모리 수요가 함께 움직입니다.", text)
            self.assertIn("키워드 그래프", text)
            self.assertGreaterEqual(len(document.inline_shapes), 2)
            style_fonts = document.styles["Normal"].element.find(qn("w:rPr")).find(qn("w:rFonts"))
            self.assertEqual(style_fonts.get(qn("w:eastAsia")), "맑은 고딕")
            found_run = False
            for paragraph in document.paragraphs:
                for run in paragraph.runs:
                    if "수출 호조" not in run.text:
                        continue
                    fonts = run._element.find(qn("w:rPr")).find(qn("w:rFonts"))
                    self.assertEqual(fonts.get(qn("w:eastAsia")), "맑은 고딕")
                    found_run = True
            self.assertTrue(found_run)

            copied = agent.copy_report(dest, folder_path / "saved")
            self.assertTrue(copied.is_file())
            self.assertEqual(copied.parent.name, "saved")

    def test_copy_requires_a_folder(self):
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / "report.docx"
            source.write_bytes(b"docx")
            with self.assertRaises(agent.NewsError):
                agent.copy_report(source, "  ")


class GraphTests(unittest.TestCase):
    def test_nodes_run_in_order_without_external_calls(self):
        drawn = agent.build_graph().get_graph()
        pairs = {(edge.source, edge.target) for edge in drawn.edges}
        self.assertIn(("__start__", "news"), pairs)
        self.assertIn(("news", "analyze"), pairs)
        self.assertIn(("analyze", "report"), pairs)
        self.assertIn(("report", "__end__"), pairs)

        calls = []

        def search(params):
            calls.append(params)
            return google_payload()

        class FakeLLM:
            def __init__(self):
                self.calls = 0
                self.messages = None
                self.schema = None

            def with_structured_output(self, schema):
                self.schema = schema
                return self

            def invoke(self, messages):
                self.calls += 1
                self.messages = messages
                return self.schema(
                    groups=[
                        agent.GroupDraft(
                            title="수출 호조",
                            summary="반도체 수출과 메모리 수요를 함께 봐야 합니다.",
                            insight="공급 변화의 시사점이 있습니다.",
                            article_numbers=[1, 2],
                        )
                    ]
                )

        fake = FakeLLM()
        with tempfile.TemporaryDirectory() as folder:
            result = agent.build_graph(search_fn=search, llm=fake).invoke(
                {"keyword": "반도체", "count": 2, "source": "google", "work_dir": folder}
            )
            self.assertEqual(len(calls), 1)
            self.assertEqual(fake.calls, 1)
            self.assertEqual(len(result["articles"]), 2)
            self.assertEqual(result["groups"][0]["제목"], "수출 호조")
            self.assertFalse(result["used_fallback"])
            self.assertTrue(Path(result["report_path"]).is_file())
            self.assertGreater(Path(result["chart_path"]).stat().st_size, 1000)
            self.assertGreater(Path(result["wordcloud_path"]).stat().st_size, 1000)
            joined = "\n".join(message.content for message in fake.messages)
            self.assertIn("반도체", joined)
            self.assertNotIn("api_key", joined.lower())


class UiTests(unittest.TestCase):
    def test_json_text_keeps_korean(self):
        text = app.articles_json(
            [{"순번": 1, "날짜": "2026-03-01", "제목": "반도체", "주요내용": "수출", "url": "https://example.com"}]
        )
        self.assertIn("반도체", text)
        self.assertNotIn("\\u", text)

    def test_screen_builds(self):
        demo = app.build_ui()
        self.assertIsNotNone(demo)
        demo.close()


if __name__ == "__main__":
    unittest.main()
